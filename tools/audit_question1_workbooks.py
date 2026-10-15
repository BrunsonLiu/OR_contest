"""Read historical workbooks without changing their values or formatting."""
import hashlib
import json
import csv
import ast
import subprocess
import sys
import math
from collections import defaultdict
from pathlib import Path

from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT))
from src.geography import warehouse_coordinates


def audit():
    records = []
    for path in sorted((ROOT/"results/question1").glob("*.xlsx")):
        workbook = load_workbook(path, read_only=True, data_only=False)
        sheets = []
        for sheet in workbook:
            rows = list(sheet.iter_rows(values_only=True))
            populated = [list(row) for row in rows if any(value is not None for value in row)]
            sheets.append({"sheet": sheet.title, "rows": sheet.max_row, "columns": sheet.max_column,
                           "first_rows": populated[:4],
                           "formula_count": sum(isinstance(value, str) and value.startswith("=") for row in rows for value in row)})
        workbook.close()
        records.append({"file": path.relative_to(ROOT).as_posix(),
                        "sha256": hashlib.sha256(path.read_bytes()).hexdigest(), "sheets": sheets})
    with (ROOT/"results/question1/Q1-result2.xlsx").open("rb") as source:
        book = load_workbook(source, read_only=True, data_only=True)
        warehouse_rows = list(book.active.values)[1:]
        book.close()
    with (ROOT/"results/question1/Q1-result3.xlsx").open("rb") as source:
        book = load_workbook(source, read_only=True, data_only=True)
        supply_rows = list(book.active.values)[1:]
        book.close()
    outbound = defaultdict(float)
    for name, customer, sku, quantity in supply_rows:
        outbound[(name, sku)] += quantity
    comparisons = [{"warehouse": name, "sku": sku, "workbook_qty": quantity,
                    "supply_outbound_qty": outbound[(name, sku)],
                    "qty_to_outbound_ratio": quantity/outbound[(name, sku)]}
                   for name, sku, quantity in warehouse_rows]
    current = ROOT/"results/experiments/audit-fixed-20261015/question1"
    csv_schemas = {}
    for path in sorted(current.glob("*.csv")):
        with path.open(encoding="utf-8-sig", newline="") as source:
            reader = csv.DictReader(source)
            csv_schemas[path.name] = {"columns": reader.fieldnames, "rows": sum(1 for _ in reader)}
    old_source = subprocess.check_output(["git", "show", "60d508d:src/visualization/visualize_problem1_images.py"], cwd=ROOT).decode("utf-8")
    old_coordinates = next(ast.literal_eval(node.value) for node in ast.parse(old_source).body
                           if isinstance(node, ast.Assign) and any(isinstance(target, ast.Name) and target.id == "warehouse_coords" for target in node.targets))
    coordinates = warehouse_coordinates()
    with (current/"仓库处置订单量文件_cplex.csv").open(encoding="utf-8-sig", newline="") as source:
        volumes = list(csv.DictReader(source))
    with (current/"仓网供应关系文件_cplex.csv").open(encoding="utf-8-sig", newline="") as source:
        current_outbound = defaultdict(float)
        for row in csv.DictReader(source):
            current_outbound[(row["Name"], row["SKU"])] += float(row["Qty（单位：吨）"])
    assert set(current_outbound) == {(row["Name"], row["SKU"]) for row in volumes}
    assert all(math.isclose(float(row["Qty（单位：吨）"]), current_outbound[(row["Name"], row["SKU"])], abs_tol=1e-7)
               for row in volumes), "当前 CSV 处理量与逐路线出库量不一致"
    coverage = {"old_coordinate_keys": len(old_coordinates), "actual_warehouse_keys": len(coordinates),
                "missing_actual_warehouses": sorted(set(coordinates)-set(old_coordinates)),
                "nonexistent_warehouse_keys": sorted(set(old_coordinates)-set(coordinates)),
                "active_warehouses_missing_before": sorted({row["Name"] for row in volumes}-set(old_coordinates)),
                "outbound_tons_omitted_before": sum(float(row["Qty（单位：吨）"]) for row in volumes if row["Name"] not in old_coordinates),
                "coordinates": coordinates}
    plot_root = ROOT/"results/experiments/q1-mapping-final-20261015"
    for filename in csv_schemas:
        assert (current/filename).read_bytes() == (plot_root/"question1"/filename).read_bytes()
    plot_runs = {}
    for name in ("visualize_problem1_images", "visualize_problem1_map", "visualize_problem1_overtime"):
        run = json.loads((plot_root/"_runs"/name/"run.json").read_text(encoding="utf-8"))
        assert run["status"] == "completed" and run["exit_code"] == 0
        plot_runs[name] = run["status"]
    plot_files = {path.name: hashlib.sha256(path.read_bytes()).hexdigest()
                  for path in sorted((plot_root/"question1").iterdir()) if path.is_file()}
    return {"workbooks": records, "historical_quantity_comparison": comparisons, "coordinate_coverage": coverage,
            "current_csv_schemas": csv_schemas, "current_outbound_reconciled_rows": len(volumes),
            "plot_evidence": {"runs": plot_runs, "files_sha256": plot_files, "csv_identical_to_corrected_experiment": True,
                              "delivery_format": "PNG only; HTML outputs removed at user request."},
            "scope": "Read-only workbook inspection; no recalculation or source-file edits."}


if __name__ == "__main__":
    result = audit()
    output = ROOT/"docs/audits/2026-10-15/question1_workbooks.json"
    output.write_text(json.dumps(result, ensure_ascii=False, indent=2, default=str), encoding="utf-8")
    print(json.dumps(result, ensure_ascii=False, indent=2, default=str))
